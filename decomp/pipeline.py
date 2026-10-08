"""Glue code: workbook -> facts/AI input (preprocess) and facts + AI reply -> HTML (render)."""
import json
import logging
from datetime import datetime
from pathlib import Path

from . import config
from . import workbook_reader as wr
from .hierarchy import build_tree, read_summary_block, NUMERIC_KEYS
from .commentary import parse_commentary, attach_commentary
from .calculations import analyse
from .output_builder import build_facts, build_ai_inputs, parse_report_date
from .html_builder import parse_ai_reply, audit_ai_text, render_html
from .utils import WorkbookFormatError, norm_label, fmt_k, VERIFIED, INFERRED

log = logging.getLogger(__name__)


def _summary_rows(ws, rules, required, anchor, start, desk, multiplier, text_keys=()):
    if start is None or start >= ws.max_row:
        return {}
    try:
        lay = wr.detect_layout(ws, rules, [k for k in required if k not in text_keys][:2], anchor,
                               search_from=start)
    except WorkbookFormatError as exc:
        log.warning("Summary table on '%s' not usable: %s", ws.title, exc)
        return {}
    if lay is None:
        return {}
    rows = wr.read_rows(ws, lay, lay.data_start, min(ws.max_row, lay.data_start + 40),
                        text_keys=text_keys, multiplier=multiplier)
    out = read_summary_block(rows, desk)
    log.info("Summary table on '%s' rows %d+: %d rows", ws.title, lay.data_start, len(out))
    return out


def preprocess(workbook_path, out_dir):
    workbook_path = Path(workbook_path)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    wb, wbf = wr.load(workbook_path)
    ws_d = wr.find_sheet(wb, config.DECOMP_SHEET_NAMES)
    ws_c = wr.find_sheet(wb, config.COMMENTARY_SHEET_NAMES)
    ignored = [s for s in wb.sheetnames if s not in (ws_d.title, ws_c.title)]
    if ignored:
        log.info("Ignoring sheets: %s", ignored)

    formula_info = {}
    for ws in (ws_d, ws_c):
        n, missing = wr.formula_summary(wbf[ws.title], ws)
        formula_info[ws.title] = {"formulas": n, "formulas_without_saved_value": missing}
        if missing:
            raise WorkbookFormatError(
                "Sheet '%s' has %d formulas with no saved result. Open the file in Excel, "
                "let it calculate, Save, and run again." % (ws.title, missing))
        if n:
            log.warning("Sheet '%s' contains %d formulas (expected hardcoded values); cached "
                        "results are used", ws.title, n)

    # ---- Decmp
    lay_d = wr.detect_layout(ws_d, config.DECOMP_COLUMN_RULES, config.DECOMP_REQUIRED,
                             config.DECOMP_ANCHOR)
    if lay_d is None:
        raise WorkbookFormatError("Could not find the decomposition header ('%s') on '%s'."
                                  % (config.DECOMP_ANCHOR, ws_d.title))
    mult_d, basis_d, why_d = wr.detect_units(ws_d, lay_d, [k for k in NUMERIC_KEYS if k in lay_d.columns])
    desk = lay_d.title or "Desk"
    rows_d = wr.read_rows(ws_d, lay_d, lay_d.data_start, ws_d.max_row, multiplier=mult_d)
    tree = build_tree(rows_d, desk)
    decmp_summary = _summary_rows(ws_d, config.DECOMP_COLUMN_RULES, config.DECOMP_REQUIRED,
                                  config.DECOMP_ANCHOR, (tree.grand_total_row or ws_d.max_row) + 1,
                                  desk, mult_d)

    # ---- Commentary
    lay_c = wr.detect_layout(ws_c, config.COMMENTARY_COLUMN_RULES, config.COMMENTARY_REQUIRED,
                             config.COMMENTARY_ANCHOR)
    if lay_c is None:
        raise WorkbookFormatError("Could not find the commentary header ('%s') on '%s'."
                                  % (config.COMMENTARY_ANCHOR, ws_c.title))
    mult_c, basis_c, why_c = wr.detect_units(ws_c, lay_c, ["actual", "estimate", "variance"])
    if mult_c != mult_d:
        tree.issues.append({"level": "WARN", "basis": INFERRED,
                            "message": "Units differ between sheets (Decmp: %s; Commentary: %s)"
                                       % (why_d, why_c)})
    if lay_c.title and norm_label(lay_c.title) != norm_label(desk):
        tree.issues.append({"level": "WARN", "basis": VERIFIED,
                            "message": "Desk title differs: Decmp '%s' vs Commentary '%s'"
                                       % (desk, lay_c.title)})
    rows_c = wr.read_rows(ws_c, lay_c, lay_c.data_start, ws_c.max_row,
                          text_keys=config.COMMENT_KEYS, multiplier=mult_c)
    from openpyxl.utils import get_column_letter
    entries, issues = parse_commentary(rows_c, desk, get_column_letter(lay_c.columns["decomp_comment"]))
    tree.issues.extend(issues)
    grand_c = next((e.excel_row for e in entries if e.norm == norm_label(desk + " total")), None)
    comm_summary = _summary_rows(ws_c, config.COMMENTARY_COLUMN_RULES, config.COMMENTARY_REQUIRED,
                                 config.COMMENTARY_ANCHOR, (grand_c or ws_c.max_row) + 1, desk,
                                 mult_c, text_keys=config.COMMENT_KEYS)
    attach_commentary(tree, entries, comm_summary, tree.issues)

    # ---- Analytics + outputs
    analyse(tree, decmp_summary, comm_summary)
    meta = {
        "source_file": workbook_path.name,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "desk": desk,
        "report_date": parse_report_date(lay_d.date_text or lay_c.date_text, workbook_path.name),
        "units": {"stored_as": "dollars" if mult_d == 1 else "thousands", "basis": basis_d,
                  "explanation": why_d, "display": "C$ thousands, whole numbers with .0"},
        "layouts": [lay_d.describe(), lay_c.describe()],
        "formulas": formula_info,
        "summary_tables_found": {"decmp": len(decmp_summary), "commentary": len(comm_summary)},
        "tolerances": {"abs_dollars": config.RECON_ABS_TOL, "pct": config.RECON_PCT_TOL},
        "ignored_sheets": ignored,
    }
    facts = build_facts(tree, meta)
    (out_dir / "facts.json").write_text(json.dumps(facts, indent=2, default=str), encoding="utf-8")
    inputs = build_ai_inputs(facts)
    for old in out_dir.glob("ai_input*.txt"):
        old.unlink()
    for name, text in inputs:
        (out_dir / name).write_text(text, encoding="utf-8")
    (out_dir / "check_report.txt").write_text(check_report(facts), encoding="utf-8")
    (out_dir / "report_draft.html").write_text(render_html(facts, {}, []), encoding="utf-8")
    return facts, [n for n, _ in inputs]


def check_report(facts):
    """Plain-text page to eyeball against Excel."""
    m = facts["meta"]
    lines = ["CHECK REPORT - %s - %s" % (m["desk"], m["source_file"]),
             "Report date: %s (%s: %s)" % (m["report_date"]["date"], m["report_date"]["basis"],
                                           m["report_date"]["note"]),
             "Units: %s (%s) - %s" % (m["units"]["stored_as"], m["units"]["basis"], m["units"]["explanation"]),
             "Formulas found: %s" % m["formulas"],
             "Summary tables matched: %s" % m["summary_tables_found"], "",
             "HIERARCHY (Actual | Estimate | Variance | recon | confidence)", "-" * 78]
    for r in facts["rows"]:
        lines.append("%s %s%-*s %14s %14s %14s  %-5s %s" % (
            r["row_id"], "  " * r["depth"], max(1, 34 - 2 * r["depth"]), r["business"][:34],
            r["actual"]["display"], r["estimate"]["display"], r["variance"]["display"],
            r["reconciliation_status"][:5], r["confidence"]))
    lines += ["", "RECONCILIATION EXCEPTIONS", "-" * 78]
    exc = [(r["business"], c) for r in facts["rows"] for c in r["reconciliation"]
           if c["status"] in ("WARN", "FAIL")]
    lines += ["%s %s - %s: diff %s, tolerance %s" % (c["status"], b, c["check"],
                                                     fmt_k(c["difference"]), fmt_k(c["tolerance"]))
              for b, c in exc] or ["none - every check within tolerance"]
    lines += ["", "FLAGS AND ISSUES", "-" * 78]
    lines += ["%s [%s] %s" % (i["level"], i["basis"], i["message"]) for i in facts["issues"]]
    lines += ["%s [%s] %s: %s" % (f["level"], f["basis"], r["business"], f["message"])
              for r in facts["rows"] for f in r["flags"]]
    return "\n".join(lines) + "\n"


def render(out_dir, ai_files=None):
    out_dir = Path(out_dir)
    facts_path = out_dir / "facts.json"
    if not facts_path.exists():
        raise WorkbookFormatError("No facts.json in %s - run the 'prepare' step first." % out_dir)
    facts = json.loads(facts_path.read_text(encoding="utf-8"))
    files = [Path(f) for f in ai_files] if ai_files else sorted(out_dir.glob("ai_response*.txt"))
    replies = {}
    for f in files:
        part = parse_ai_reply(f.read_text(encoding="utf-8"))
        log.info("Read %d AI items from %s", len(part), f.name)
        replies.update(part)
    if not files:
        log.warning("No ai_response*.txt found in %s - report will show placeholders", out_dir)
    problems = audit_ai_text(facts, replies)
    (out_dir / "ai_check.txt").write_text(
        "AI OUTPUT CHECK\n" + ("\n".join(problems) if problems else "No problems found.") + "\n",
        encoding="utf-8")
    html_path = out_dir / "report.html"
    html_path.write_text(render_html(facts, replies, problems), encoding="utf-8")
    return html_path, problems

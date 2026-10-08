"""
Step 1 - read the workbook without assuming cell addresses.

* finds sheets by name (tolerant of spacing/case, with a fuzzy fallback)
* finds the header block by an anchor phrase and maps columns by header text
* reads every row below the header as a RawRow (label + numbers + text)
* works out whether cells hold dollars or thousands
"""
import difflib
import logging
import re
import statistics
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from . import config
from .utils import (VERIFIED, INFERRED, WorkbookFormatError, clean_text,
                    to_number)

log = logging.getLogger(__name__)


@dataclass
class SheetLayout:
    sheet_name: str
    label_col: int
    columns: dict                 # key -> column index (1-based)
    header_rows: tuple            # (first, last)
    data_start: int
    title: str = ""
    date_text: str = ""

    def describe(self):
        cols = {k: get_column_letter(c) for k, c in self.columns.items()}
        return {"sheet": self.sheet_name, "label_column": get_column_letter(self.label_col),
                "columns": cols, "header_rows": list(self.header_rows),
                "data_starts_row": self.data_start, "title": self.title,
                "date_text": self.date_text}


@dataclass
class RawRow:
    sheet: str
    excel_row: int
    label: str
    values: dict = field(default_factory=dict)    # numeric key -> float (dollars) or None
    texts: dict = field(default_factory=dict)     # text key -> str
    problems: list = field(default_factory=list)
    hidden: bool = False

    @property
    def has_numbers(self):
        return any(v is not None for v in self.values.values())

    @property
    def has_text(self):
        return any(t for t in self.texts.values())


# ------------------------------------------------------------------ load ----
def load(path):
    """Open the workbook twice: cached values (what Excel shows) and formulas."""
    path = Path(path)
    if not path.exists():
        raise WorkbookFormatError("File not found: %s" % path)
    if path.suffix.lower() == ".xls":
        raise WorkbookFormatError("Old .xls format is not supported. Open it in Excel "
                                  "and 'Save As' .xlsx, then run again.")
    try:
        wb_values = load_workbook(path, data_only=True)
        wb_formulas = load_workbook(path, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise WorkbookFormatError("Could not open %s as an Excel workbook (%s). Is it "
                                  "open/locked, password protected or corrupt?" % (path.name, exc))
    log.info("Opened %s - sheets: %s", path.name, wb_values.sheetnames)
    return wb_values, wb_formulas


def find_sheet(wb, candidates, required=True):
    def n(s):
        return re.sub(r"\s+", " ", s.strip().lower())
    by_norm = {n(name): name for name in wb.sheetnames}
    for cand in candidates:
        if n(cand) in by_norm:
            return wb[by_norm[n(cand)]]
    close = difflib.get_close_matches(n(candidates[0]), list(by_norm), n=1, cutoff=0.8)
    if close:
        actual = by_norm[close[0]]
        log.warning("Sheet '%s' not found; using close match '%s'", candidates[0], actual)
        return wb[actual]
    if required:
        raise WorkbookFormatError("Required sheet %s not found. Sheets present: %s. "
                                  "Update the sheet names in config.py."
                                  % (candidates, wb.sheetnames))
    return None


def formula_summary(ws_formulas, ws_values):
    """Count formulas and check Excel saved their results (cached values)."""
    n_formulas, missing = 0, 0
    for row in ws_formulas.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                n_formulas += 1
                if ws_values[cell.coordinate].value is None:
                    missing += 1
    return n_formulas, missing


# ---------------------------------------------------------------- layout ----
def _cell_text(ws, r, c):
    v = ws.cell(row=r, column=c).value
    return clean_text(v) if isinstance(v, str) else ""


def _row_has_number(ws, r):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return True
    return False


def _merged_bottom(ws, r, c):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
            return rng.max_row
    return r


def detect_layout(ws, rules, required, anchor, search_from=1, max_scan=80):
    """Locate the header block below `search_from` and map columns by header text.

    Returns SheetLayout, or None if the anchor phrase is not found.
    """
    last = min(ws.max_row, search_from + max_scan)
    anchor_row = None
    for r in range(search_from, last + 1):
        for c in range(1, ws.max_column + 1):
            if anchor in _cell_text(ws, r, c).lower():
                anchor_row = r
                break
        if anchor_row:
            break
    if anchor_row is None:
        return None

    # Header block: anchor row, plus up to 2 text-only rows above it, plus
    # rows covered by merged header cells / text-only header rows below it.
    first = anchor_row
    for r in (anchor_row - 1, anchor_row - 2):
        if r >= search_from and not _row_has_number(ws, r):
            first = r
        else:
            break
    last_hdr = anchor_row
    for c in range(1, ws.max_column + 1):
        for r in range(first, anchor_row + 1):
            if ws.cell(row=r, column=c).value is not None:
                last_hdr = max(last_hdr, _merged_bottom(ws, r, c))
    while True:  # a following row with text in 2+ columns and no numbers is still header
        r = last_hdr + 1
        texts = sum(1 for c in range(1, ws.max_column + 1) if _cell_text(ws, r, c))
        if r <= ws.max_row and texts >= 2 and not _row_has_number(ws, r):
            last_hdr = r
        else:
            break

    headers = {}
    for c in range(1, ws.max_column + 1):
        parts = [_cell_text(ws, r, c) for r in range(first, last_hdr + 1)]
        txt = re.sub(r"\s+", " ", " ".join(p for p in parts if p)).strip().lower()
        if txt:
            headers[c] = txt

    columns, claimed = {}, set()
    for key, mode, phrases in rules:
        for c in sorted(headers):
            if c in claimed:
                continue
            h = headers[c]
            hit = (any(p in h for p in phrases) if mode == "contains"
                   else any(h == p for p in phrases))
            if hit:
                columns[key] = c
                claimed.add(c)
                break
    missing = [k for k in required if k not in columns]
    if missing:
        raise WorkbookFormatError(
            "Sheet '%s': could not find column(s) %s in header rows %d-%d. Headers seen: %s. "
            "Update the *_COLUMN_RULES in config.py." % (ws.title, missing, first, last_hdr,
                                                         {get_column_letter(c): h for c, h in headers.items()}))
    for key, c in columns.items():
        if ws.column_dimensions[get_column_letter(c)].hidden:
            log.warning("Sheet '%s': column %s (%s) is HIDDEN but was mapped", ws.title,
                        get_column_letter(c), key)

    # Label column: the column (left of the first mapped column) with most text below.
    leftmost = min(columns.values())
    best, best_count = None, -1
    for c in range(1, leftmost):
        count = sum(1 for r in range(last_hdr + 1, min(ws.max_row, last_hdr + 60) + 1)
                    if _cell_text(ws, r, c))
        if count > best_count:
            best, best_count = c, count
    if not best or best_count == 0:
        raise WorkbookFormatError("Sheet '%s': could not find the business-name column "
                                  "left of the data columns." % ws.title)

    title, date_text = "", ""
    for r in range(first, last_hdr + 1):
        t = _cell_text(ws, r, best)
        if t.lower().startswith("daily"):
            date_text = t
        elif t and not title:
            title = t
    layout = SheetLayout(ws.title, best, columns, (first, last_hdr), last_hdr + 1, title, date_text)
    log.info("Layout %s", layout.describe())
    return layout


# ------------------------------------------------------------------ rows ----
def read_rows(ws, layout, start, end, text_keys=(), multiplier=1.0):
    rows = []
    for r in range(start, end + 1):
        label = clean_text(ws.cell(row=r, column=layout.label_col).value)
        rr = RawRow(ws.title, r, label, hidden=bool(ws.row_dimensions[r].hidden))
        for key, c in layout.columns.items():
            raw = ws.cell(row=r, column=c).value
            if key in text_keys:
                rr.texts[key] = "" if raw is None else str(raw)
                continue
            num, problem = to_number(raw)
            rr.values[key] = None if num is None else num * multiplier
            if problem:
                rr.problems.append("%s%d: %s" % (get_column_letter(c), r, problem))
        rows.append(rr)
    return rows


# ----------------------------------------------------------------- units ----
_SCALED = re.compile(r"[0#?](,+)(?=[^0#?]*$)")


def _is_thousand_scaled(fmt):
    for section in (fmt or "").split(";"):
        s = re.sub(r'"[^"]*"|\[[^\]]*\]|_.|\\.', "", section)
        if _SCALED.search(s):
            return True
    return False


def detect_units(ws, layout, numeric_keys):
    """Return (multiplier_to_dollars, basis, explanation)."""
    mode = config.UNITS_MODE.lower()
    if mode == "dollars":
        return 1.0, VERIFIED, "config UNITS_MODE='dollars'"
    if mode == "thousands":
        return 1000.0, VERIFIED, "config UNITS_MODE='thousands'"
    scaled, total, mags = 0, 0, []
    for r in range(layout.data_start, ws.max_row + 1):
        for key in numeric_keys:
            c = layout.columns.get(key)
            if not c:
                continue
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                total += 1
                scaled += _is_thousand_scaled(cell.number_format)
                if cell.value:
                    mags.append(abs(cell.value))
    if total == 0:
        raise WorkbookFormatError("Sheet '%s' has no numeric values under the header." % ws.title)
    median = statistics.median(mags) if mags else 0
    if scaled / total > 0.5:
        return 1.0, VERIFIED, ("%d/%d number formats divide by 1,000 for display, so cells "
                               "hold dollars (median |value| %.0f)" % (scaled, total, median))
    if median >= 5000:
        return 1.0, INFERRED, ("number formats do not scale, but median |value| is %.0f, "
                               "so cells are assumed to hold dollars" % median)
    return 1000.0, INFERRED, ("median |value| is %.1f, so cells are assumed to hold "
                              "thousands and are multiplied by 1,000" % median)
